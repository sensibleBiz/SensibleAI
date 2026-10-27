from __future__ import annotations

import json
import os
import re
import sys
from collections import defaultdict
from pathlib import Path

_deps = Path(os.environ.get("LOCALAPPDATA", "")) / "Temp" / "codex_doc_deps"
if _deps.exists():
    sys.path.insert(0, str(_deps))

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
WORK = ROOT / ".codex-work"
OUTPUT_DIR = ROOT / "artifacts"
OUTPUT = OUTPUT_DIR / "Sensible_CRM_Mobile_API_Reference.docx"

NAVY = "163A5F"
PALE_BLUE = "EAF2F8"
PALE_GRAY = "F5F7F9"
MID_GRAY = "6B7280"
LIGHT_BORDER = "D9D9D9"
BLACK = "000000"
WHITE = "FFFFFF"
GREEN = "18794E"
ORANGE = "A24E00"
RED = "A12828"


def load_json(path: Path):
    return json.loads(path.read_text(encoding="utf-8"))


inventory = load_json(WORK / "api-inventory.json")
schemas = load_json(WORK / "api-schemas.json")


def norm_path(value: str) -> str:
    value = value.replace("\\", "/")
    if value != "/" and value.endswith("/"):
        value = value[:-1]
    return value


for endpoint in inventory:
    endpoint["fullPath"] = norm_path(endpoint["fullPath"])
    endpoint["routeFile"] = endpoint["routeFile"].replace("\\", "/")
    if endpoint.get("controllerFile"):
        endpoint["controllerFile"] = endpoint["controllerFile"].replace("\\", "/")
    for schema in (endpoint.get("requestSchemas") or {}).values():
        schema["file"] = schema["file"].replace("\\", "/")


def set_font(run, name="Arial", size=None, bold=None, italic=None, color=BLACK):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    return run


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=LIGHT_BORDER, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_cell_width(cell, inches):
    cell.width = Inches(inches)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def set_paragraph_spacing(paragraph, before=0, after=0, line=1.0, keep_with_next=None):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    if keep_with_next is not None:
        fmt.keep_with_next = keep_with_next


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rel_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), NAVY)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), "Arial")
    r_fonts.set(qn("w:hAnsi"), "Arial")
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "20")
    r_pr.extend([r_fonts, color, underline, size])
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.extend([r_pr, text_node])
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_field(paragraph, instruction, display=""):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = display
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    set_font(run, size=9, color=MID_GRAY)


def add_code_block(doc, text, size=8.5):
    paragraph = doc.add_paragraph(style="Code Block")
    set_paragraph_spacing(paragraph, before=2, after=7, line=1.05)
    run = paragraph.add_run(text)
    set_font(run, name="Consolas", size=size, color=BLACK)
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), PALE_GRAY)
    p_pr.append(shd)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "120")
    ind.set(qn("w:right"), "120")
    p_pr.append(ind)
    return paragraph


def add_bullet(doc, text, level=0):
    paragraph = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    set_paragraph_spacing(paragraph, after=2, line=1.08)
    set_font(paragraph.add_run(text), size=10.2)
    return paragraph


def add_number(doc, text):
    paragraph = doc.add_paragraph(style="List Number")
    set_paragraph_spacing(paragraph, after=3, line=1.08)
    set_font(paragraph.add_run(text), size=10.2)
    return paragraph


def add_label_paragraph(doc, label, text):
    paragraph = doc.add_paragraph()
    set_paragraph_spacing(paragraph, after=5, line=1.1)
    set_font(paragraph.add_run(label), size=10.3, bold=True)
    set_font(paragraph.add_run(text), size=10.3)
    return paragraph


def add_heading(doc, text, level=1):
    paragraph = doc.add_heading(text, level=level)
    set_paragraph_spacing(paragraph, before=12 if level == 1 else 9, after=5, line=1.0, keep_with_next=True)
    return paragraph


def add_table(doc, headers, widths, rows, font_size=8.8, header_fill=NAVY):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    header = table.rows[0]
    set_repeat_table_header(header)
    prevent_row_split(header)
    for idx, (cell, label, width) in enumerate(zip(header.cells, headers, widths)):
        set_cell_width(cell, width)
        set_cell_shading(cell, header_fill)
        set_cell_margins(cell, top=100, bottom=100, start=110, end=110)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
        set_paragraph_spacing(paragraph, after=0, line=1.0)
        set_font(paragraph.add_run(label), size=9, bold=True, color=WHITE)
    for row_index, values in enumerate(rows):
        row = table.add_row()
        if row_index % 2 == 1:
            for cell in row.cells:
                set_cell_shading(cell, PALE_BLUE)
        for col_index, (cell, value, width) in enumerate(zip(row.cells, values, widths)):
            set_cell_width(cell, width)
            set_cell_margins(cell, top=90, bottom=90, start=100, end=100)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            set_paragraph_spacing(paragraph, after=0, line=1.03)
            if isinstance(value, list):
                for run_idx, spec in enumerate(value):
                    if run_idx:
                        paragraph.add_run("\n")
                    run = paragraph.add_run(spec.get("text", ""))
                    set_font(
                        run,
                        name=spec.get("font", "Arial"),
                        size=spec.get("size", font_size),
                        bold=spec.get("bold"),
                        italic=spec.get("italic"),
                        color=spec.get("color", BLACK),
                    )
            else:
                set_font(paragraph.add_run(str(value)), size=font_size)
    spacer = doc.add_paragraph()
    set_paragraph_spacing(spacer, after=2)
    return table


def human_words(value: str) -> str:
    value = re.sub(r"([a-z0-9])([A-Z])", r"\1 \2", value)
    value = re.sub(r"[_-]+", " ", value)
    return " ".join(word.capitalize() for word in value.split())


def enum_value(value):
    if isinstance(value, str):
        return json.dumps(value)
    if value is None:
        return "null"
    if value is True:
        return "true"
    if value is False:
        return "false"
    return str(value)


def desc_signature(desc, depth=0):
    if not isinstance(desc, dict):
        return "unknown"
    kind = desc.get("type", "unknown")
    if kind == "object":
        parts = []
        for name, child in (desc.get("fields") or {}).items():
            optional = "?" if child.get("optional") else ""
            parts.append(f"{name}{optional}: {desc_signature(child, depth + 1)}")
        base = "{ " + "; ".join(parts) + " }"
    elif kind == "array":
        base = f"Array<{desc_signature(desc.get('item') or {}, depth + 1)}>"
    elif kind == "record":
        base = f"Record<string, {desc_signature(desc.get('item') or {}, depth + 1)}>"
    elif kind == "enum":
        base = " | ".join(enum_value(v) for v in (desc.get("enum") or [])) or "enum"
    elif kind == "literal":
        base = enum_value((desc.get("enum") or [None])[0])
    elif kind == "union":
        base = " | ".join(desc_signature(v, depth + 1) for v in (desc.get("options") or []))
    elif kind == "intersection":
        base = " & ".join(desc_signature(v, depth + 1) for v in (desc.get("options") or []))
    elif kind == "tuple":
        base = "[" + ", ".join(desc_signature(v, depth + 1) for v in (desc.get("options") or [])) + "]"
    elif kind == "date":
        base = "Date"
    elif kind in {"any", "unknown", "never", "null", "string", "number", "boolean"}:
        base = kind
    else:
        base = kind
    constraints = list(desc.get("constraints") or [])
    if desc.get("nullable"):
        base += " | null"
    if "default" in desc:
        constraints.append(f"default {enum_value(desc['default'])}")
    constraints = [c for c in constraints if c != "no unknown fields"]
    if constraints:
        base += " [" + ", ".join(constraints) + "]"
    return base


def auth_label(endpoint):
    middleware = " ".join(endpoint.get("middleware") or [])
    path = endpoint["fullPath"]
    if "requireEventGatewayKey" in middleware:
        return "Event key: X-Event-Gateway-Key"
    if "requireBiometricIngestAuth" in middleware:
        return "Biometric bearer secret"
    if "requirePostgresAuth" in middleware:
        suffix = " + X-Postgres-Session: 1" if "POSTGRES_SESSION_HEADER" in middleware else ""
        return "PostgreSQL bearer" + suffix
    if re.search(r"\brequireAuth\b", middleware):
        return "Firebase bearer"
    if path in {"/api/v1/auth/refresh", "/api/v1/auth/logout"}:
        return "Refresh cookie + X-Postgres-Session: 1"
    if path == "/api/v1/biometric/transactions":
        return "Device integration key in body"
    return "Public"


def request_label(endpoint):
    parts = []
    for part in ("params", "query", "body"):
        schema = (endpoint.get("requestSchemas") or {}).get(part)
        if schema:
            parts.append(f"{part}: {schema['name']}")
    if parts:
        return "; ".join(parts)
    path = endpoint["fullPath"]
    if path == "/api/v1/biometric/transactions":
        return "body: { integrationKey; trans[] }"
    if path in {"/api/v1/auth/refresh", "/api/v1/auth/logout"}:
        return "HttpOnly refresh cookie"
    return "No validated payload"


def response_display_type(response):
    value = response.get("bodyType") or response.get("body") or "unknown"
    value = re.sub(r"\s+", " ", value).strip()
    body = response.get("body") or ""
    if re.search(r"success\s*:\s*true", body):
        value = re.sub(r"success:\s*boolean", "success: true", value, count=1)
    elif re.search(r"success\s*:\s*false", body):
        value = re.sub(r"success:\s*boolean", "success: false", value, count=1)
    return value


response_models = []
response_model_index = {}
response_used_by = defaultdict(list)
endpoint_response_refs = {}
for endpoint in inventory:
    refs = []
    seen = set()
    for response in endpoint.get("response") or []:
        display_type = response_display_type(response)
        key = display_type
        if key not in response_model_index:
            model_id = f"R{len(response_models) + 1:03d}"
            response_model_index[key] = model_id
            response_models.append((model_id, display_type))
        model_id = response_model_index[key]
        status = str(response.get("status") or "200")
        pair = (status, model_id)
        if pair not in seen:
            refs.append(pair)
            seen.add(pair)
        response_used_by[model_id].append(f"{endpoint['method']} {endpoint['fullPath']}")
    endpoint_response_refs[id(endpoint)] = refs


GROUPS = [
    ("service", "Service status", "Root and health endpoints used for liveness, readiness, and server identification."),
    ("auth", "Authentication onboarding and sessions", "Password login, refresh-cookie lifecycle, Firebase-assisted onboarding, QR sign-in, and current-user context."),
    ("assistant", "CRM assistant", "Authenticated natural-language CRM queries and confirmed employee-target updates."),
    ("business", "Business configuration", "Business model, terminology, pipeline stages, and custom field configuration."),
    ("customers", "Customers", "Customer creation, search, matching, detail, and enquiry history."),
    ("leads", "Leads reports and sales targets", "Lead records, assignment and stage changes, sharing, dashboards, funnel reports, repeated enquiries, and employee targets."),
    ("activities", "Activities", "Lead-linked and global CRM activity management."),
    ("calls", "Calls", "Call creation, employee and lead call lists, detail, and reports."),
    ("attendance", "Attendance and biometric devices", "Daily and monthly attendance, event ingestion, deletion, PostgreSQL biometric ingestion, and MINOP-compatible device callbacks."),
    ("leave", "Leave balances and holidays", "Leave types, balances, balance history, holiday calendar, requests, and approval status."),
    ("team", "Team tree", "Employee profiles, memberships, departments, roles, zones, and reporting relationships."),
    ("tasks", "Tasks", "Task creation, filters, optimistic-concurrency updates, status transitions, history, and soft deletion."),
    ("products", "Products and categories", "Product and category CRUD operations and product search."),
    ("warehouse", "Warehouses and inventory", "Warehouse records, stock, ledger entries, adjustments, and transfer workflows."),
    ("expense", "Expenses and budgets", "Budget allocation, transaction history, expense creation, approval, history, and deletion."),
    ("payroll", "Payroll", "Employee payroll data, rules, preview, salary structure, and payroll runs."),
    ("access", "Access control", "Permissions, user access, manageable outlets, and feature configuration."),
    ("outlets", "Outlets and plans", "Outlet discovery and maintenance plus the plan feature catalog."),
    ("integrations", "Integrations", "Integration records and Facebook page connections."),
    ("quotations", "Quotations", "Quotation creation and the server-side WhatsApp PDF proxy."),
    ("tracking", "User tracking", "Location point ingestion and date-filtered tracking history."),
    ("whatsapp", "WhatsApp chats", "Conversation persistence, provider events, sending, and read state."),
    ("automation", "WhatsApp automation", "Template-backed automation rules and test sends."),
    ("ai", "AI contact jobs", "AI contact-generation job lifecycle."),
    ("generated", "Generated contact imports", "Spreadsheet import records, generated contacts, reassignment, and removal."),
    ("public", "Public links and quick contacts", "Public event lead capture and authenticated public-link management."),
    ("reference", "Reference geography", "Country, state, city, and zone lookup data."),
    ("events", "Event and payment gateway", "Event registration, transaction lookup, PhonePe payment initiation, and status checks guarded by a shared event key."),
    ("internal", "Internal diagnostics", "Authenticated Data Connect smoke-check endpoint; do not expose it as a normal mobile feature."),
]


def group_key(endpoint):
    route_file = endpoint["routeFile"].lower()
    path = endpoint["fullPath"]
    if route_file.endswith("app.ts") or "health.routes" in route_file:
        return "service"
    if any(x in route_file for x in ("auth.routes", "session.routes", "currentuser.routes")):
        return "auth"
    if "/assistant/" in path:
        return "assistant"
    if "businessconfiguration" in route_file:
        return "business"
    if "customers" in route_file:
        return "customers"
    if "leads" in route_file:
        return "leads"
    if "activities" in route_file:
        return "activities"
    if "calls" in route_file:
        return "calls"
    if "attendance" in route_file or "/biometric/" in path:
        return "attendance"
    if "/leave" in path:
        return "leave"
    if "team" in route_file:
        return "team"
    if "tasks" in route_file:
        return "tasks"
    if "products" in route_file:
        return "products"
    if "warehouse" in route_file:
        return "warehouse"
    if "expense" in route_file:
        return "expense"
    if "payroll" in route_file:
        return "payroll"
    if "accesscontrol" in route_file:
        return "access"
    if "outlets" in route_file:
        return "outlets"
    if "integrations" in route_file:
        return "integrations"
    if "quotations" in route_file:
        return "quotations"
    if "usertracking" in route_file:
        return "tracking"
    if "/whatsapp/" in path:
        return "whatsapp"
    if "whatsappautomation" in route_file:
        return "automation"
    if "aicontacts" in route_file:
        return "ai"
    if "generatedcontacts" in route_file:
        return "generated"
    if "publiclinks" in route_file:
        return "public"
    if "referencedata" in route_file:
        return "reference"
    if "events" in route_file:
        return "events"
    if "dataconnectsmoke" in route_file:
        return "internal"
    return "internal"


grouped = defaultdict(list)
leave_aliases = []
for endpoint in inventory:
    if endpoint["fullPath"].startswith("/api/v1/leave"):
        leave_aliases.append(endpoint)
        continue
    grouped[group_key(endpoint)].append(endpoint)


def make_document():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.68)
    section.bottom_margin = Inches(0.62)
    section.left_margin = Inches(0.66)
    section.right_margin = Inches(0.66)
    section.header_distance = Inches(0.25)
    section.footer_distance = Inches(0.28)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10.3)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.1

    title = styles["Title"]
    title.font.name = "Arial"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    title.font.size = Pt(30)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(BLACK)
    title.paragraph_format.space_after = Pt(10)

    for level, size in ((1, 18), (2, 14), (3, 11.5)):
        style = styles[f"Heading {level}"]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.paragraph_format.keep_with_next = True

    if "Code Block" not in styles:
        code_style = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    else:
        code_style = styles["Code Block"]
    code_style.font.name = "Consolas"
    code_style._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
    code_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
    code_style.font.size = Pt(8.5)
    code_style.font.color.rgb = RGBColor.from_string(BLACK)

    # Minimal running header/footer are appropriate for this long reference manual.
    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_paragraph_spacing(hp, after=0)
    set_font(hp.add_run("Sensible CRM Mobile API Reference"), size=8.5, color=MID_GRAY)
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(fp, after=0)
    set_font(fp.add_run("Page "), size=8.5, color=MID_GRAY)
    add_field(fp, "PAGE", "1")

    doc.core_properties.title = "Sensible CRM Mobile API Reference"
    doc.core_properties.subject = "Server endpoints request contracts and response bodies"
    doc.core_properties.keywords = "Sensible CRM, mobile API, REST, endpoints, request schema, response body"

    # Cover page
    spacer = doc.add_paragraph()
    set_paragraph_spacing(spacer, before=78, after=0)
    title_p = doc.add_paragraph(style="Title")
    set_font(title_p.add_run("Sensible CRM Mobile API Reference"), size=30, bold=True)
    subtitle = doc.add_paragraph()
    set_paragraph_spacing(subtitle, after=18, line=1.0)
    set_font(subtitle.add_run("Server endpoints request contracts and response bodies"), size=15, color=MID_GRAY)
    intro = doc.add_paragraph()
    set_paragraph_spacing(intro, after=12, line=1.18)
    set_font(
        intro.add_run(
            "This document is the mobile integration contract for the CRM-Web backend. "
            "It lists every mounted Express route, the required authentication mode, the request validator, "
            "the success or handled response body, and the source location."
        ),
        size=11.5,
    )
    meta_rows = [
        ("Code snapshot", "7 October 2026"),
        ("Mounted URL and method combinations", "196"),
        ("Primary routes", "183"),
        ("Legacy leave aliases", "13"),
        ("Unmounted legacy route declarations", "3"),
    ]
    add_table(doc, ["Document scope", "Value"], [2.65, 4.55], meta_rows, font_size=10)
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=9, after=0)
    set_font(p.add_run("Recommended production API base"), size=10.5, bold=True)
    link_p = doc.add_paragraph()
    set_paragraph_spacing(link_p, after=0)
    add_hyperlink(
        link_p,
        "https://sensible-crm-api-700529442570.asia-south1.run.app/api/v1",
        "https://sensible-crm-api-700529442570.asia-south1.run.app/api/v1",
    )
    doc.add_page_break()

    # Contents
    add_heading(doc, "Contents", 1)
    toc = doc.add_paragraph()
    set_paragraph_spacing(toc, after=8)
    set_font(toc.add_run("Use Word's Navigation Pane to jump between the headings listed below."), size=10.2, color=MID_GRAY)
    add_heading(doc, "Document map", 2)
    for text in (
        "Integration settings explains server URLs, headers, cookies, authentication, rate limits, and common errors.",
        "Endpoint catalog lists all primary mounted routes and the 13 leave aliases.",
        "Request validation index identifies every Zod validator used by a mounted route and its source file.",
        "Response schema reference gives the exact TypeScript response-body shape inferred from each controller response.",
        "Deployment and external service appendix separates CRM endpoints from upstream and third-party URLs.",
    ):
        add_bullet(doc, text)
    doc.add_page_break()

    # Integration settings
    add_heading(doc, "Integration settings", 1)
    opening = doc.add_paragraph()
    set_paragraph_spacing(opening, after=7, line=1.13)
    set_font(
        opening.add_run(
            "Use one explicit base URL in the mobile build configuration. The code contains separate production, test, and local values; "
            "do not infer the release target from a development environment file. All application endpoints are under /api/v1 except the root server ping at /."
        ),
        size=10.5,
    )
    add_heading(doc, "Server links", 2)
    server_rows = [
        (
            "Production fallback in web client",
            "https://sensible-crm-api-700529442570.asia-south1.run.app/api/v1",
            "Recommended mobile production base. Defined in src/shared/api/apiClient.ts and eventApi.ts.",
        ),
        (
            "Development test value",
            "https://sensible-crm-api-sqlconnect-test-jd772w3xnq-el.a.run.app/api/v1",
            "Defined in .env.development. Treat as test or legacy until deployment ownership confirms it.",
        ),
        (
            "Local backend",
            "http://localhost:3001/api/v1",
            "Defined by backend PORT=3001 and the local Vite target. A physical phone must use the computer LAN address, not localhost.",
        ),
        (
            "Local web proxy",
            "/api/v1",
            "The Vite web client uses same-origin proxying in development. Native mobile apps need an explicit absolute URL.",
        ),
    ]
    add_table(doc, ["Environment", "Base URL", "Mobile guidance"], [1.55, 2.75, 2.9], server_rows, font_size=8.8)
    add_label_paragraph(
        doc,
        "Release check. ",
        "The PostgreSQL session client requires an explicit HTTPS URL in production. Set the mobile environment URL deliberately and verify /health/ready before release.",
    )

    add_heading(doc, "Authentication and headers", 2)
    auth_rows = [
        ("PostgreSQL user route", "Authorization: Bearer <accessToken>\nX-Postgres-Session: 1", "Use for normal CRM features. Sending the selector header globally is safe."),
        ("Refresh and logout", "X-Postgres-Session: 1\nCookie: sensible_crm_refresh=<HttpOnly value>", "Use a cookie-aware HTTP client. The cookie path is /api/v1."),
        ("Firebase onboarding route", "Authorization: Bearer <Firebase ID token>", "Used by QR claim and the Firebase-assisted onboarding status/business flow."),
        ("Event gateway", "X-Event-Gateway-Key: <shared secret>", "Kiosk or payment-callback credential. Do not package the secret in a general-user app."),
        ("PostgreSQL biometric ingest", "Authorization: Bearer <BIOMETRIC_INGEST_TOKEN>", "Dedicated device or trusted integration credential."),
        ("JSON body", "Content-Type: application/json\nAccept: application/json", "Request body limit is 1 MiB."),
        ("Tracing", "X-Request-Id: <optional UUID or trace ID>", "The server forwards or creates an ID and returns it in the response."),
    ]
    add_table(doc, ["Route class", "Required headers", "Notes"], [1.65, 2.6, 2.95], auth_rows, font_size=8.7)

    add_heading(doc, "Session lifecycle", 2)
    for step in (
        "Call POST /api/v1/auth/login with identifier and password. Store only the returned access token in secure device storage; the refresh token is set as an HttpOnly cookie.",
        "Attach Authorization: Bearer <accessToken> and X-Postgres-Session: 1 to PostgreSQL CRM requests.",
        "On the first 401, make one synchronized POST /api/v1/auth/refresh request with the cookie, replace the access token, and retry the original request once.",
        "If refresh fails, clear local session state and require sign-in. Do not run parallel refresh calls for multiple failed requests.",
        "Call POST /api/v1/auth/logout when the user signs out so the server revokes the refresh session and clears the cookie.",
    ):
        add_number(doc, step)
    add_heading(doc, "Login example", 3)
    add_code_block(
        doc,
        "POST {API_BASE}/auth/login\n"
        "X-Postgres-Session: 1\n"
        "Content-Type: application/json\n\n"
        '{\n  "identifier": "employee@example.com",\n  "password": "••••••••"\n}',
    )
    add_code_block(
        doc,
        '{\n  "success": true,\n  "data": {\n'
        '    "accessToken": "<JWT>",\n    "accessTokenExpiresIn": 900,\n'
        '    "user": {"id":"<uuid>","email":"employee@example.com","phone":null,"displayName":"Employee"},\n'
        '    "context": {"tenantId":"<uuid>","isTenantAdmin":false,"outletIds":["<uuid>"]}\n'
        "  }\n}\n\nSet-Cookie: sensible_crm_refresh=<HttpOnly refresh token>",
    )

    add_heading(doc, "Response envelopes and errors", 2)
    add_label_paragraph(doc, "Standard success. ", "Most controllers return the following envelope. The endpoint catalog points to the exact response schema for the data value.")
    add_code_block(doc, '{ "success": true, "data": <endpoint-specific value> }')
    add_label_paragraph(doc, "Standard failure. ", "Validation, authentication, authorization, not-found, conflict, and server errors normally use this shape.")
    add_code_block(doc, '{ "success": false, "message": "Human-readable message", "code": "MACHINE_READABLE_CODE" }')
    error_rows = [
        ("400", "VALIDATION_ERROR or INVALID_JSON", "Invalid request field or malformed JSON."),
        ("401", "INVALID_AUTH_TOKEN or AUTHENTICATION_REQUIRED", "Access token is absent, invalid, expired, or no longer maps to an active membership."),
        ("403", "SESSION_HEADER_REQUIRED or domain permission code", "Required selector header or domain permission is missing."),
        ("404", "ROUTE_NOT_FOUND or domain not-found code", "Unknown path or missing entity."),
        ("409", "Domain conflict code", "Version, state transition, duplicate, or QR-session conflict."),
        ("429", "RATE_LIMIT_EXCEEDED", "Retry after the Retry-After interval."),
        ("500", "INTERNAL_SERVER_ERROR", "Unexpected server error. Production omits the stack trace."),
        ("502 or 503", "Provider or configuration code", "External provider failure or a required server integration is disabled."),
    ]
    add_table(doc, ["HTTP", "Typical code", "Meaning"], [0.8, 2.55, 3.85], error_rows, font_size=8.9)
    add_label_paragraph(doc, "Rate limits. ", "The general /api/v1 limiter permits 300 requests per minute per IP. /api/v1/auth also has a 20 requests per minute per IP limiter. Read X-RateLimit-Limit, X-RateLimit-Remaining, X-RateLimit-Reset, and Retry-After.")
    add_label_paragraph(doc, "Timestamps and pagination. ", "Many CRM timestamps are epoch milliseconds. Dates use YYYY-MM-DD where a validator says date. Cursors are opaque strings and must be returned unchanged.")
    doc.add_page_break()

    # Endpoint catalog
    add_heading(doc, "Endpoint catalog", 1)
    intro = doc.add_paragraph()
    set_paragraph_spacing(intro, after=7)
    set_font(
        intro.add_run(
            "Every mounted primary route appears below. Request entries name the exact Zod validator; response entries point to the full body schema in the Response schema reference. "
            "Thrown domain errors use the common failure envelope described above."
        ),
        size=10.4,
    )
    summary_rows = []
    for key, title_text, _summary in GROUPS:
        count = len(grouped.get(key, []))
        if count:
            summary_rows.append((title_text, str(count)))
    summary_rows.append(("Legacy /api/v1/leave aliases", str(len(leave_aliases))))
    add_table(doc, ["API area", "Routes"], [6.3, 0.9], summary_rows, font_size=9.2)

    for key, title_text, summary in GROUPS:
        endpoints = sorted(grouped.get(key, []), key=lambda e: (e["fullPath"], e["method"]))
        if not endpoints:
            continue
        add_heading(doc, title_text, 2)
        paragraph = doc.add_paragraph()
        set_paragraph_spacing(paragraph, after=5)
        set_font(paragraph.add_run(summary), size=10)
        rows = []
        for endpoint in endpoints:
            source = f"{endpoint['routeFile']}:{endpoint['line']}"
            path_cell = [
                {"text": endpoint["fullPath"], "font": "Consolas", "size": 8.4, "bold": True},
                {"text": source, "font": "Arial", "size": 7.4, "color": MID_GRAY},
            ]
            access_request = [
                {"text": auth_label(endpoint), "size": 8.4, "bold": True},
                {"text": request_label(endpoint), "font": "Consolas", "size": 7.7, "color": MID_GRAY},
            ]
            refs = endpoint_response_refs[id(endpoint)]
            response_lines = []
            for status, model_id in refs:
                status_label = "dynamic" if not re.fullmatch(r"\d+", status) else status
                response_lines.append({"text": f"{status_label}  {model_id}", "font": "Consolas", "size": 8.2, "bold": True})
            rows.append((
                [{"text": endpoint["method"], "size": 8.2, "bold": True, "color": NAVY}],
                path_cell,
                access_request,
                response_lines,
            ))
        add_table(doc, ["Method", "Path and source", "Access and request", "Response"], [0.65, 2.65, 2.65, 1.25], rows, font_size=8.2)
        if key == "leave":
            add_heading(doc, "Legacy leave aliases", 3)
            p = doc.add_paragraph()
            set_paragraph_spacing(p, after=5)
            set_font(p.add_run("The following /api/v1/leave paths are mounted aliases of /api/v1/attendance/leaves. Use the attendance path in new mobile code."), size=9.8)
            alias_rows = []
            for endpoint in sorted(leave_aliases, key=lambda e: (e["fullPath"], e["method"])):
                canonical = endpoint["fullPath"].replace("/api/v1/leave", "/api/v1/attendance/leaves", 1)
                alias_rows.append((endpoint["method"], endpoint["fullPath"], canonical))
            add_table(doc, ["Method", "Legacy path", "Preferred path"], [0.7, 3.25, 3.25], alias_rows, font_size=8.2)

    doc.add_page_break()

    # Compact request-validator index. Full field contracts remain in the named source files,
    # while every endpoint above identifies exactly which schema applies to params, query, and body.
    add_heading(doc, "Request validation index", 1)
    p = doc.add_paragraph()
    set_paragraph_spacing(p, after=7)
    set_font(
        p.add_run(
            "The endpoint catalog names every Zod validator used for path parameters, query strings, and JSON bodies. "
            "This index groups those validator names by source file for quick lookup; custom refinements and permission checks can add domain rules after structural validation."
        ),
        size=10.3,
    )
    schemas_by_file = defaultdict(set)
    for endpoint in inventory:
        for schema in (endpoint.get("requestSchemas") or {}).values():
            schemas_by_file[schema["file"]].add(schema["name"])
    schema_index_rows = []
    for file_name in sorted(schemas_by_file):
        schema_index_rows.append((
            [{"text": file_name, "font": "Consolas", "size": 7.9}],
            [{"text": ", ".join(sorted(schemas_by_file[file_name])), "font": "Consolas", "size": 7.9}],
        ))
    add_table(doc, ["Validator source", "Schemas used by mounted routes"], [3.15, 4.05], schema_index_rows, font_size=8.0)

    doc.add_page_break()

    # Response schema reference
    add_heading(doc, "Response schema reference", 1)
    p = doc.add_paragraph()
    set_paragraph_spacing(p, after=7)
    set_font(
        p.add_run(
            "Each response ID is the body passed to response.json or response.send by the controller. The shapes are inferred from the TypeScript program, so named interfaces such as CustomerListResult and PostgresCurrentAccess remain named where the controller returns them directly."
        ),
        size=10.3,
    )
    for start in range(0, len(response_models), 25):
        chunk = response_models[start:start + 25]
        add_heading(doc, f"Response schemas {chunk[0][0]} to {chunk[-1][0]}", 2)
        rows = []
        for model_id, display_type in chunk:
            used = list(dict.fromkeys(response_used_by[model_id]))
            used_text = "; ".join(used[:4]) + (f"; plus {len(used)-4} more" if len(used) > 4 else "")
            contract = [
                {"text": display_type, "font": "Consolas", "size": 7.8},
                {"text": "Used by " + used_text, "size": 7.3, "color": MID_GRAY},
            ]
            rows.append((
                [{"text": model_id, "font": "Consolas", "size": 8.4, "bold": True, "color": NAVY}],
                contract,
            ))
        add_table(doc, ["ID", "Response body contract"], [0.75, 6.45], rows, font_size=8.0)

    doc.add_page_break()

    # Deployment appendix and exclusions
    add_heading(doc, "Deployment and external service appendix", 1)
    add_heading(doc, "Unmounted legacy route declarations", 2)
    p = doc.add_paragraph()
    set_paragraph_spacing(p, after=5)
    set_font(
        p.add_run(
            "backend/src/modules/attendance/legacyAttendanceGateway.routes.ts declares three routes, but the router is not imported or mounted by backend/src/routes/index.ts. They therefore have no callable server URL in the current application."
        ),
        size=10.2,
    )
    legacy_rows = [
        ("POST", "/attendance-report", "Unmounted legacy attendance report handler"),
        ("POST", "/webhook", "Unmounted legacy device webhook"),
        ("GET", "/api/filtered-logs", "Unmounted legacy filtered-log handler"),
    ]
    add_table(doc, ["Method", "Declared path", "Status"], [0.75, 2.6, 3.85], legacy_rows, font_size=9)

    add_heading(doc, "External and upstream URLs found in source", 2)
    p = doc.add_paragraph()
    set_paragraph_spacing(p, after=5)
    set_font(p.add_run("These URLs are not mounted CRM REST endpoints. Most must remain server-side because their credentials or provider contracts are not suitable for a general mobile client."), size=10.2)
    external_rows = [
        ("WhatsApp template create and get", "https://create-whatsapp-template-700529442570.asia-south1.run.app", "Server only"),
        ("MSG91 WhatsApp bulk and direct messaging", "Configured WHATSAPP_API_URL and WHATSAPP_MESSAGE_API_URL", "Server only"),
        ("Quotation PDF WhatsApp proxy target", "https://asia-south1-sensiblebizpro.cloudfunctions.net/msg91Whatsapp/send_pdf_whatsapp", "Server only"),
        ("Event invoice service", "https://event-invoice-401022448274.asia-south1.run.app/helloHttp", "Server only"),
        ("PhonePe offline test and production hosts", "Configured in backend event settings", "Server only"),
        ("Biomax device API", "http://www.biomaxpay.com/api/MobileApp", "Trusted integration only"),
        ("Sendesk ticket lookup", "https://sendesk-401022448274.asia-south1.run.app/issue/getWithFilters", "Direct web-client dependency; external to CRM API"),
        ("Sendesk ticket update", "https://sendesk-401022448274.asia-south1.run.app/issue/updateTicket", "Direct web-client dependency; external to CRM API"),
    ]
    add_table(doc, ["Purpose", "URL or configuration", "Mobile use"], [2.0, 3.4, 1.8], external_rows, font_size=8.4)

    add_heading(doc, "Mobile release checklist", 2)
    checklist = [
        "Set an explicit HTTPS production API base and remove local or test overrides from the release build.",
        "Store access tokens in secure storage and never log tokens, cookies, event keys, biometric secrets, or provider responses that contain credentials.",
        "Use a cookie jar that persists the HttpOnly refresh cookie when persistent sign-in is required.",
        "Send X-Postgres-Session: 1 for PostgreSQL session and CRM requests.",
        "Synchronize token refresh and retry an unauthorized request only once.",
        "Send epoch timestamps in milliseconds and preserve pagination cursors exactly.",
        "Honor task version and outlet-feature expectedRevision conflicts instead of overwriting newer server state.",
        "Keep event gateway, biometric ingest, WhatsApp, PhonePe, and other provider secrets outside a general-user mobile binary.",
        "Exercise /health/live and /health/ready against the selected environment before publishing.",
    ]
    for item in checklist:
        add_bullet(doc, "[ ] " + item)

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    result = make_document()
    print(result)
